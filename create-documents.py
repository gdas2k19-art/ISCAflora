from pathlib import Path
import base64, re, shutil, json, markdown
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/home/user')
SITE = ROOT / 'isca-website'
DOWNLOADS = SITE / 'public/downloads'
REPORT = ROOT / 'research/ISCA-Flora-Research-and-Website-Copy.md'

text = REPORT.read_text().replace('not a independently', 'not an independently').replace('Eight authentic still frames from the official Annual Function reel:', 'Eight authentic still frames prepared from the official Annual Function reel, with selected frames integrated:')
REPORT.write_text(text)
shutil.copy(REPORT, DOWNLOADS / 'research-and-website-copy.md')

def data(path, mime):
    return f'data:{mime};base64,' + base64.b64encode(Path(path).read_bytes()).decode()

logo = data(SITE / 'public/media/isca-logo-small.png', 'image/png')
flora_logo = data(SITE / 'public/media/flora-logo.webp', 'image/webp')
award = data(SITE / 'public/media/annual-award.webp', 'image/webp')
flora = data(SITE / 'public/media/flora-performance.webp', 'image/webp')
fonts = ''
for family, file, weight in [('Outfit','outfit-variable.woff2',700),('Plus Jakarta Sans','plus-jakarta-sans-variable.woff2',400),('Plus Jakarta Sans','plus-jakarta-sans-variable.woff2',600)]:
    fonts += f"@font-face{{font-family:'{family}';font-weight:{weight};src:url({data(SITE / ('public/fonts/'+file),'font/woff2')}) format('woff2');}}"

report_css = fonts + '''
*{box-sizing:border-box}body{font-family:'Plus Jakarta Sans',sans-serif;color:#334155;margin:0;font-size:10.5pt;line-height:1.6}h1,h2,h3{font-family:Outfit,sans-serif;line-height:1.3;break-after:avoid}h1{color:#243fc5;font-size:29pt;margin:8mm 0 2mm;letter-spacing:-1pt}h2{font-size:17pt;color:#243fc5;letter-spacing:-.4pt;margin:9mm 0 4mm}h3{font-size:12pt;color:#1e4237;margin:6mm 0 3mm}p{margin:0 0 3mm;orphans:3;widows:3}ul,ol{margin:2mm 0 4mm;padding-left:6mm}li{margin-bottom:1.7mm}a{color:#2944bc;text-decoration:none;overflow-wrap:anywhere}strong{color:#1c2a43;font-weight:600}table{width:100%;border-collapse:collapse;font-size:8.4pt;line-height:1.5;margin:5mm 0 6mm;table-layout:fixed}thead{display:table-header-group}th{background:#edf0fa;color:#243fc5;font-weight:600;text-align:left}th,td{border:1px solid #dfe5e7;padding:2.5mm 2.6mm;vertical-align:top;overflow-wrap:anywhere}tr{break-inside:avoid}hr{border:0;border-top:1px solid #e4e9df;margin:7mm 0}blockquote{border-left:3px solid #459573;margin:5mm 0;padding:3mm 5mm;background:#f1f6ec;color:#4e6253}blockquote p:last-child{margin:0}code{font-size:9pt;color:#385045}.brand-top{display:flex;align-items:center;gap:5mm;border-bottom:2px solid #243fc5;padding-bottom:4mm;margin-bottom:5mm}.brand-top img{width:17mm;height:16mm;object-fit:contain}.brand-top strong{font-family:Outfit;font-size:16pt;color:#17233a}.brand-top span{display:block;font-size:8pt;color:#617076;margin-top:1mm}.document-label{margin-left:auto;font-size:8pt;letter-spacing:1pt;color:#65705f;text-align:right}.owner-note{padding:4mm;background:#f1f5e8;border:1px solid #dce5cd;border-radius:2mm;font-size:9pt;margin-top:3mm;color:#566549}
'''
report_html = f'''<!doctype html><html><head><meta charset="utf-8"><title>ISCA & Flora — Research and Website Copy</title><style>{report_css}</style></head><body><div class="brand-top"><img src="{logo}" alt="ISCA original emblem"><div><strong>ISCA & Flora Play School</strong><span>Education with morality · Topsia, Kolkata</span></div><div class="document-label">RESEARCH & WEBSITE COPY<br>12 OCTOBER 2026</div></div>{markdown.markdown(text, extensions=['tables','fenced_code'])}</body></html>'''
(ROOT / 'research/research-print.html').write_text(report_html)

# Editable OOXML write-up, with source links and tables preserved.
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)
sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.72)
normal = doc.styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string('334155')
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.2
for name, size in [('Title',29),('Heading 1',18),('Heading 2',14),('Heading 3',11.5)]:
    st=doc.styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('243FC5');st.paragraph_format.space_before=Pt(16);st.paragraph_format.space_after=Pt(8)
header = sec.header.paragraphs[0]
header.add_run().add_picture(str(SITE/'public/media/isca-logo-small.png'),width=Inches(.38))
run=header.add_run('   ISCA & FLORA PLAY SCHOOL  /  RESEARCH & WEBSITE COPY');run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('667580')
footer=sec.footer.paragraphs[0];footer.alignment=2
run=footer.add_run('12 October 2026  ·  Owner approval required  ·  ');run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('667580')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def add_link(paragraph, label, url):
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'243FC5');props.append(color);r.append(props)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);paragraph._p.append(h)

def rich(p, value):
    parts=re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\)|https?://\S+)',value)
    for piece in parts:
        if not piece:continue
        if piece.startswith('**') and piece.endswith('**'):p.add_run(piece[2:-2]).bold=True
        elif re.fullmatch(r'\[[^\]]+\]\(https?://[^)]+\)',piece):
            m=re.match(r'\[([^\]]+)\]\(([^)]+)\)',piece);add_link(p,m[1],m[2])
        elif piece.startswith('http'):add_link(p,piece,piece)
        else:p.add_run(piece.replace('`',''))

lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line or line=='---':i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=lines[i].strip();i+=1
            if re.fullmatch(r'[|:\-\s]+',row):continue
            rows.append([c.strip() for c in row.strip('|').split('|')])
        if rows:
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Table Grid'
            for n,cell in enumerate(table.rows[0].cells):
                rich(cell.paragraphs[0],rows[0][n]);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF0FA');cell._tc.get_or_add_tcPr().append(sh)
                for rr in cell.paragraphs[0].runs:rr.bold=True;rr.font.size=Pt(9)
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows[1:]:
                cells=table.add_row().cells
                for n,v in enumerate(row):
                    rich(cells[n].paragraphs[0],v)
                    for rr in cells[n].paragraphs[0].runs:rr.font.size=Pt(8.5)
            doc.add_paragraph()
        continue
    if line.startswith('# '):p=doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        if line.startswith('## 6.') or line.startswith('## 10.'):doc.add_page_break()
        doc.add_heading(line[3:],level=1)
    elif line.startswith('### '):doc.add_heading(line[4:],level=2)
    elif line.startswith('- '):rich(doc.add_paragraph(style='List Bullet'),line[2:])
    elif re.match(r'^\d+\. ',line):rich(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',line))
    elif line.startswith('> '):rich(doc.add_paragraph(style='Quote'),line[2:])
    else:
        # Join soft line breaks without merging new blocks.
        buffer=line
        while i+1<len(lines) and lines[i+1].strip() and not re.match(r'^(#|- |\||>|---|\d+\. )',lines[i+1].strip()):
            i+=1;buffer+=' '+lines[i].strip()
        rich(doc.add_paragraph(),buffer)
    i+=1
file=ROOT/'ISCA-Flora-Research-and-Website-Copy.docx';doc.save(file)

brochure_css = fonts + '''
@page{size:A4;margin:0}*{box-sizing:border-box}body{margin:0;font-family:'Plus Jakarta Sans',sans-serif;font-size:10pt;color:#334155}a{color:#243fc5;text-decoration:none}.paper{width:210mm;height:297mm;padding:16mm 17mm;position:relative;overflow:hidden;break-after:page;background:#fcfcf7}.paper:last-child{break-after:auto}.brand{display:flex;align-items:center;gap:4mm;border-bottom:1px solid #e1e6dc;padding-bottom:5mm}.brand img{width:17mm;height:16mm;object-fit:contain}.brand strong{font-family:Outfit;font-size:19pt;color:#17233a}.brand span{display:block;font-size:8pt;color:#65716b;margin-top:1mm}.tag{margin-left:auto;font-size:7pt;color:#60715a;text-align:right;line-height:1.7}.kicker{font-size:7pt;letter-spacing:1.6pt;color:#63765a;margin:9mm 0 4mm}h1,h2,h3{font-family:Outfit;line-height:1.15;margin:0;color:#17233a}h1{font-size:36pt;letter-spacing:-1.6pt;margin-bottom:5mm}h1 span,h2 span{color:#243fc5}h2{font-size:27pt;letter-spacing:-1pt}h3{font-size:13pt;letter-spacing:-.3pt}p{font-size:9.5pt;line-height:1.65;margin:0 0 4mm;color:#63716d}.intro{max-width:147mm}.photo{width:100%;height:58mm;object-fit:cover;object-position:50% 35%;border-radius:4mm;margin:5mm 0 2mm}.photo-note{font-size:6.5pt;color:#697762;margin-bottom:6mm}.cards{display:grid;grid-template-columns:1fr 1fr;gap:4mm}.card{padding:4mm;border:1px solid #e0e5db;border-radius:3mm;background:white}.card:nth-child(1){background:#eff2fc}.card:nth-child(2){background:#faf1eb}.card:nth-child(3){background:#f7f4e8}.card:nth-child(4){background:#eef5ea}.card .label{font-size:6.5pt;letter-spacing:.9pt;color:#667b5d;margin-bottom:2mm}.card p{font-size:8pt;line-height:1.6;margin:2mm 0 0}.section-note{font-size:7pt;line-height:1.7;color:#66765f;margin-top:5mm}.footer{position:absolute;bottom:13mm;left:17mm;right:17mm;border-top:1px solid #dce3d6;padding-top:3mm;display:flex;justify-content:space-between;font-size:7pt;color:#5c6d5b}.flora-page{background:#eaf3e5}.flora-page h2 span{color:#347650}.flora-header{display:flex;align-items:center;gap:4mm;margin:7mm 0 5mm}.flora-header img{width:20mm;height:20mm;border-radius:3mm}.flora-header strong{font-family:Outfit;font-size:16pt;color:#2e623f}.flora-header span{display:block;font-size:7pt;color:#567248;margin-top:2mm}.flora-page .photo{height:54mm;object-position:center;margin-top:4mm}.ages{display:grid;grid-template-columns:repeat(3,1fr);gap:3mm;margin:6mm 0}.age{background:#ffffffc9;border:1px solid #cfdec4;border-radius:3mm;padding:4mm}.age strong{font-size:10pt;color:#2f653f}.age .years{display:block;font-size:8pt;color:#507742;margin:3mm 0}.age p{font-size:7.5pt;line-height:1.6;margin:0}.contact{border-top:1px solid #cbdcc2;margin-top:6mm;padding-top:5mm}.contact h3{font-size:17pt;color:#2d5f3d;margin-bottom:4mm}.contact p{font-size:9pt;color:#4e6849}.contact .numbers{font-size:13pt;color:#214f35;font-weight:600;line-height:1.7;margin:2mm 0 4mm}.fine{font-size:7pt;line-height:1.7;color:#5f7555;margin:4mm 0}.link-row{display:flex;gap:6mm;align-items:center;font-size:8pt;margin-top:4mm}
'''
brochure_css += '''
.paper:not(.flora-page) .photo{height:52mm;object-position:center 18%}
.flora-header{margin:5mm 0 3mm}.flora-header img{width:16mm;height:16mm}
.flora-page .photo{height:42mm;object-position:center 35%}.ages{margin:5mm 0}.contact{margin-top:4mm;padding-top:4mm}
.contact h3{font-size:15pt;margin-bottom:3mm}.contact p{font-size:8.5pt;line-height:1.5}.contact .numbers{font-size:12pt;margin:1mm 0 3mm}
.contact .fine{font-size:6.5pt;line-height:1.5;margin:3mm 0}.link-row{margin-top:3mm;font-size:7.5pt}
'''
brochure = f'''<!doctype html><html><head><meta charset="utf-8"><title>ISCA & Flora — Programme Overview</title><style>{brochure_css}</style></head><body>
<section class="paper"><header class="brand"><img src="{logo}" alt="ISCA"><div><strong>ISCA<span style="display:inline;color:#df6653">.</span></strong><span>The Institute of Science, Commerce & Arts</span></div><div class="tag">PROGRAMME OVERVIEW<br>TOPSIA, KOLKATA · SINCE 2012</div></header><div class="kicker">EDUCATION WITH MORALITY</div><h1>Big dreams.<br>Bright minds.<br><span>Better futures.</span></h1><p class="intro">A learning family in Topsia, Kolkata. Explore your next academic step at ISCA — and an early-years beginning at Flora Play School.</p><img class="photo" src="{award}" alt="Student recognition at ISCA’s annual function"><div class="photo-note">Real school moment · Official ISCA Annual Function reel · September 2026</div><div class="cards"><div class="card"><div class="label">CLASSES VII–XII</div><h3>School learning</h3><p>Enquire about Science, Commerce and Arts subject support. Public material lists CBSE, ICSE / ISC and WBBSE / WBCHSE.</p></div><div class="card"><div class="label">JEE & NEET</div><h3>Bigger ambitions</h3><p>Discuss current entrance-exam preparation options, supported subjects, faculty and available batches.</p></div><div class="card"><div class="label">B.COM & B.A</div><h3>Your next chapter</h3><p>Ask about undergraduate subject support for your semester and syllabus. No degree-awarding claim is made.</p></div><div class="card"><div class="label">HISTORICALLY ADVERTISED TRACKS</div><h3>Skills for what’s next</h3><p>Ask whether spoken English, basic computer learning, and Java / Python / C / C++ courses are currently available.</p></div></div><p class="section-note">Please confirm current courses, fees, schedules and eligibility directly. Listed school boards refer to tuition / exam support, not affiliation. This overview is based on public material reviewed on 12 October 2026.</p><footer class="footer"><span>26 Topsia Road · Beside Belal Masjid · Kolkata 700039</span><span>01 / 02</span></footer></section>
<section class="paper flora-page"><header class="brand"><img src="{logo}" alt="ISCA"><div><strong>One learning family.</strong><span>ISCA & Flora Play School</span></div><div class="tag">EARLY-YEARS BEGINNINGS<br>AGE BANDS TO BE CONFIRMED AT ENQUIRY</div></header><div class="flora-header"><img src="{flora_logo}" alt="Flora’s official logo"><div><strong>Flora Play School</strong><span>Where every child blossoms</span></div></div><h2>Small steps.<br><span>Wonderful beginnings.</span></h2><p style="margin-top:4mm">A nurturing introduction to learning, friendship and self-expression. A place for little discoveries — and a whole lot of growing.</p><img class="photo" src="{flora}" alt="Flora’s little learners on the annual function stage"><div class="photo-note">Flora children at ISCA’s Annual Function · Still from the official September 2026 event reel</div><div class="ages"><div class="age"><strong>Montessori 1</strong><span class="years">2.5–3.5 years</span><p>A happy first step into exploration, early friendships and everyday discoveries.</p></div><div class="age"><strong>Montessori 2</strong><span class="years">3.5–4.5 years</span><p>Growing independence, self-expression and readiness for more structured learning.</p></div><div class="age"><strong>L.K.G / KG</strong><span class="years">4.5–5.5 years</span><p>An early-years pathway towards foundational learning and school readiness.</p></div></div><p class="fine">Age bands follow Flora’s March 2026 public Facebook information. Confirm cut-off dates, eligibility and current admissions with the school. No historical seasonal offer is presented as current.</p><div class="contact"><h3>Let’s talk about your tomorrow.</h3><p>26 Topsia Road, beside Belal Masjid · Kolkata, West Bengal 700039</p><div class="numbers"><a href="tel:+917003561315">+91 70035 61315</a> &nbsp; / &nbsp; <a href="tel:+919883345594">+91 98833 45594</a></div><p>ISCA: <a href="mailto:iscawb@gmail.com">iscawb@gmail.com</a><br>Flora: <a href="mailto:floraplayschool@gmail.com">floraplayschool@gmail.com</a></p><div class="link-row"><a href="https://maps.app.goo.gl/1enz3x4GEpiizJkp6">Open Google Maps ↗</a><a href="https://www.instagram.com/floraplayschool/">Meet Flora on Instagram ↗</a><a href="https://www.youtube.com/@iscakolkata">ISCA on YouTube ↗</a></div><p class="fine">Please call ahead to arrange a school visit or demo class. The team will confirm timing and availability.</p></div><footer class="footer"><span>Education with morality · Every learner. Every possibility.</span><span>02 / 02</span></footer></section>
</body></html>'''
(ROOT / 'research/programme-overview-print.html').write_text(brochure)
print('Created editable DOCX and print-ready HTML documents.')
